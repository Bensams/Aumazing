"""Builds the content-validity calculator that goes with the expert rating
sheet (build_rating_sheet.js), from the same templates.json.

    python docs/validation/src/build_calculator.py

Output: docs/validation/Aumazing_Content_Validity_Calculator.xlsx

Every result is a live formula over the yellow input cells, so the
workbook recalculates as ratings are typed in.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
OUT = HERE.parent / "Aumazing_Content_Validity_Calculator.xlsx"
T = json.loads((HERE / "templates.json").read_text(encoding="utf-8"))

EXPERTS = 10
FONT = "Times New Roman"
INPUT_FILL = PatternFill("solid", start_color="FFFF00")
INPUT_FONT = Font(name=FONT, size=11, color="0000FF")
BODY = Font(name=FONT, size=11)
BOLD = Font(name=FONT, size=11, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

domains = T["home"]["domains"]
home = [(f"A{i + 1}", "A. Home", domains[x["domain"]], x["text"])
        for i, x in enumerate(T["home"]["items"])]
after = [(f"B{i + 1}", "B. After the session",
          "Each game" if q["scope"] == "game" else "Once", q["text"])
         for i, q in enumerate(T["observation"]["items"])]
before = [("C1", "C. Before the session", "", "Who is watching?")] + [
    (f"C{i + 2}", "C. Before the session", "", q["text"])
    for i, q in enumerate(T["intake"]["items"])]
flags = [
    (f"F{n}", "B. Flags", q["text"], f'"{o["label"]}" shows: {o["flag"]}')
    for n, (q, o) in enumerate(
        ((q, o) for q in T["observation"]["items"]
         for o in q.get("options", []) if o.get("flag")), start=1)]


def style_header(ws, row, headers, widths):
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = BOX
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=5)


def rating_sheet(wb, title, heading, note, items, parts):
    """One row per item: expert ratings in, I-CVI and modified kappa out."""
    ws = wb.create_sheet(title)
    ws["A1"] = heading
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].font = BODY
    ws["A3"] = ("Enter each expert's rating (1-4) in the yellow cells; leave a "
                "cell empty if that expert did not rate the item.")
    ws["A3"].font = BODY
    hr = 5
    first_e = 5  # column E
    last_e = first_e + EXPERTS - 1
    ecol = lambda i: get_column_letter(first_e + i)  # noqa: E731
    cN, cA, cI, cPc, cK, cT, cD, cKi = (
        get_column_letter(last_e + k) for k in range(1, 9))
    headers = (["No.", "Part", "Area / asked", "Question"]
               + [f"Expert {i + 1}" for i in range(EXPERTS)]
               + ["Experts (N)", "Rated 3 or 4 (A)", "I-CVI", "Pc",
                  "Modified kappa (k*)", "Required I-CVI", "Decision",
                  "k* rating"])
    widths = [6, 18, 16, 48] + [9] * EXPERTS + [10, 10, 9, 9, 11, 10, 11, 11]
    style_header(ws, hr, headers, widths)

    dv = DataValidation(type="whole", operator="between", formula1="1",
                        formula2="4", allow_blank=True,
                        error="Enter a whole number from 1 to 4.",
                        errorTitle="Rating")
    ws.add_data_validation(dv)

    rows = {}
    for r, (no, part, area, text) in enumerate(items, start=hr + 1):
        rows.setdefault(part, []).append(r)
        for c, v in enumerate([no, part, area, text], start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY
            cell.alignment = CENTER if c == 1 else WRAP
            cell.border = BOX
        for i in range(EXPERTS):
            cell = ws[f"{ecol(i)}{r}"]
            cell.fill = INPUT_FILL
            cell.font = INPUT_FONT
            cell.alignment = CENTER
            cell.border = BOX
            dv.add(cell)
        rng = f"{ecol(0)}{r}:{ecol(EXPERTS - 1)}{r}"
        formulas = {
            cN: f"=COUNT({rng})",
            cA: f'=COUNTIF({rng},">=3")',
            cI: f'=IF({cN}{r}=0,"",{cA}{r}/{cN}{r})',
            # Probability of chance agreement (Polit, Beck & Owen, 2007).
            cPc: f'=IF({cN}{r}=0,"",COMBIN({cN}{r},{cA}{r})*0.5^{cN}{r})',
            cK: f'=IF({cN}{r}=0,"",({cI}{r}-{cPc}{r})/(1-{cPc}{r}))',
            # Lynn (1986): all must agree with 3-5 experts; .78 with 6+.
            cT: f'=IF({cN}{r}=0,"",IF({cN}{r}<6,1,0.78))',
            cD: (f'=IF({cN}{r}=0,"",IF({cN}{r}<3,"Need 3+ experts",'
                 f'IF({cI}{r}>={cT}{r},"Keep","Revise")))'),
            cKi: (f'=IF({cN}{r}=0,"",IF({cK}{r}>0.74,"Excellent",'
                  f'IF({cK}{r}>=0.6,"Good",IF({cK}{r}>=0.4,"Fair","Poor"))))'),
        }
        for col, f in formulas.items():
            cell = ws[f"{col}{r}"]
            cell.value = f
            cell.font = BODY
            cell.alignment = CENTER
            cell.border = BOX
            if col in (cI, cPc, cK, cT):
                cell.number_format = "0.00"
    last_row = hr + len(items)
    ws.conditional_formatting.add(
        f"{cD}{hr + 1}:{cD}{last_row}",
        CellIsRule(operator="equal", formula=['"Revise"'],
                   font=Font(name=FONT, bold=True, color="C00000")))

    # Scale-level summary for each part.
    r = last_row + 2
    ws.cell(row=r, column=4, value="Scale-level results").font = BOLD
    for c, h in enumerate(["Part", "Items rated", "S-CVI/Ave",
                           "S-CVI/UA", "Items to revise"], start=4):
        cell = ws.cell(row=r + 1, column=c, value=h)
        cell.font = BOLD
        cell.alignment = CENTER
        cell.border = BOX
    for k, part in enumerate(parts):
        rr = r + 2 + k
        lo, hi = rows[part][0], rows[part][-1]
        irng = f"{cI}{lo}:{cI}{hi}"
        vals = [
            part,
            f"=COUNT({irng})",
            f'=IF(COUNT({irng})=0,"",AVERAGE({irng}))',
            f'=IF(COUNT({irng})=0,"",COUNTIF({irng},1)/COUNT({irng}))',
            f'=COUNTIF({cD}{lo}:{cD}{hi},"Revise")',
        ]
        for c, v in enumerate(vals, start=4):
            cell = ws.cell(row=rr, column=c, value=v)
            cell.font = BODY
            cell.border = BOX
            cell.alignment = WRAP if c == 4 else CENTER
            if c in (6, 7):
                cell.number_format = "0.00"
    ws.cell(row=r + 2 + len(parts), column=4, value=(
        "S-CVI/Ave: average I-CVI; .90 or higher is excellent (Polit & Beck, "
        "2006). S-CVI/UA: share of items every expert rated 3 or 4.")).font = BODY
    ws.cell(row=r + 2 + len(parts), column=4).alignment = WRAP
    ws.cell(row=hr, column=last_e + 3).comment = Comment(
        "I-CVI = A / N: the proportion of experts rating the item 3 or 4.",
        "Aumazing")
    return ws


def yes_no_sheet(wb, items):
    ws = wb.create_sheet("Yes-No Checks")
    ws["A1"] = "Yes / No judgements"
    ws["A1"].font = TITLE
    ws["A2"] = ("Part A: is the statement in the right area? Parts B and C: are "
                "the answer choices complete and appropriate? Flags: keep the "
                "flag? Enter Y or N in the yellow cells.")
    ws["A2"].font = BODY
    hr = 4
    first_e = 5
    ecol = lambda i: get_column_letter(first_e + i)  # noqa: E731
    cN = get_column_letter(first_e + EXPERTS)
    cY = get_column_letter(first_e + EXPERTS + 1)
    cP = get_column_letter(first_e + EXPERTS + 2)
    cD = get_column_letter(first_e + EXPERTS + 3)
    headers = (["No.", "Check", "Question", "Item"]
               + [f"Expert {i + 1}" for i in range(EXPERTS)]
               + ["Answered", "Yes", "% Yes", "Decision"])
    style_header(ws, hr, headers, [6, 26, 30, 44] + [9] * EXPERTS + [10, 8, 9, 11])
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True,
                        error="Enter Y or N.", errorTitle="Yes or No")
    ws.add_data_validation(dv)
    for r, (no, check, q, item) in enumerate(items, start=hr + 1):
        for c, v in enumerate([no, check, q, item], start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BODY
            cell.alignment = CENTER if c == 1 else WRAP
            cell.border = BOX
        for i in range(EXPERTS):
            cell = ws[f"{ecol(i)}{r}"]
            cell.fill = INPUT_FILL
            cell.font = INPUT_FONT
            cell.alignment = CENTER
            cell.border = BOX
            dv.add(cell)
        rng = f"{ecol(0)}{r}:{ecol(EXPERTS - 1)}{r}"
        for col, f in {
            cN: f"=COUNTA({rng})",
            cY: f'=COUNTIF({rng},"Y")',
            cP: f'=IF({cN}{r}=0,"",{cY}{r}/{cN}{r})',
            cD: f'=IF({cN}{r}=0,"",IF({cP}{r}>=0.78,"OK","Review"))',
        }.items():
            cell = ws[f"{col}{r}"]
            cell.value = f
            cell.font = BODY
            cell.alignment = CENTER
            cell.border = BOX
            if col == cP:
                cell.number_format = "0%"
    ws.cell(row=hr + len(items) + 2, column=4, value=(
        "Items answered Yes by fewer than 78% of experts are reviewed with the "
        "experts' comments, using the same cut-off as the I-CVI.")).font = BODY
    return ws


def read_me(wb):
    ws = wb.active
    ws.title = "Read Me"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    lines = [
        ("Aumazing content validity calculator", None, TITLE),
        ("Purpose", "Computes content validity from the experts' ratings on the "
         "Aumazing Expert Content Validity Rating Sheet. Enter each returned "
         "sheet as one expert column.", None),
        ("Where to type", "Only the yellow cells (blue numbers): 1-4 on "
         "Relevance and Clarity, Y or N on Yes-No Checks. Everything else is a "
         "formula.", None),
        ("Experts", "Up to 10. Lynn (1986) recommends at least 3; with 3-5 "
         "experts every expert must rate an item 3 or 4, with 6 or more the "
         "I-CVI must be at least .78.", None),
        ("I-CVI", "Experts rating the item 3 or 4, divided by experts who "
         "rated it.", None),
        ("Modified kappa (k*)", "I-CVI adjusted for chance agreement. Pc = "
         "COMBIN(N, A) x 0.5^N; k* = (I-CVI - Pc) / (1 - Pc). Above .74 "
         "excellent, .60-.74 good, .40-.59 fair (Polit, Beck, & Owen, 2007).",
         None),
        ("S-CVI/Ave", "Average I-CVI of a part; .90 or higher is excellent "
         "(Polit & Beck, 2006).", None),
        ("Clarity", "Computed the same way to guide rewording; it is not a "
         "validity index.", None),
        ("Next step", "Revise items marked Revise using the comments and have "
         "them rated again. Enter the final version in the practitioner portal "
         "and mark it validated.", None),
        ("Source of the items", "The app's built-in questions, exported from "
         "packages/questionnaire_core.", None),
        ("", "", None),
        ("References", "Lynn, M. R. (1986). Determination and quantification of "
         "content validity. Nursing Research, 35(6), 382-385.", BOLD),
        ("", "Polit, D. F., & Beck, C. T. (2006). The content validity index: "
         "Are you sure you know what's being reported? Critique and "
         "recommendations. Research in Nursing & Health, 29(5), 489-497.", None),
        ("", "Polit, D. F., Beck, C. T., & Owen, S. V. (2007). Is the CVI an "
         "acceptable indicator of content validity? Appraisal and "
         "recommendations. Research in Nursing & Health, 30(4), 459-467.", None),
    ]
    for r, (a, b, font) in enumerate(lines, start=1):
        ws.cell(row=r, column=1, value=a).font = font or BOLD
        if b is not None:
            cell = ws.cell(row=r, column=2, value=b)
            cell.font = BODY
            cell.alignment = WRAP

    # A worked example with live formulas: five experts rate one item.
    r0 = len(lines) + 2
    ws.cell(row=r0, column=1, value="Worked example").font = BOLD
    ws.cell(row=r0, column=2, value=(
        "Five experts rate one item 4, 3, 4, 2, 4. Change the yellow cells to "
        "see the results move.")).font = BODY
    for i, v in enumerate([4, 3, 4, 2, 4]):
        cell = ws.cell(row=r0 + 1 + i, column=1, value=f"Expert {i + 1}")
        cell.font = BODY
        cell = ws.cell(row=r0 + 1 + i, column=2, value=v)
        cell.fill = INPUT_FILL
        cell.font = INPUT_FONT
        cell.alignment = Alignment(horizontal="left")
    rng = f"B{r0 + 1}:B{r0 + 5}"
    ex = [
        ("N", f"=COUNT({rng})"),
        ("A (rated 3 or 4)", f'=COUNTIF({rng},">=3")'),
        ("I-CVI", f"=B{r0 + 7}/B{r0 + 6}"),
        ("Pc", f"=COMBIN(B{r0 + 6},B{r0 + 7})*0.5^B{r0 + 6}"),
        ("k*", f"=(B{r0 + 8}-B{r0 + 9})/(1-B{r0 + 9})"),
        ("Decision", f'=IF(B{r0 + 8}>=IF(B{r0 + 6}<6,1,0.78),"Keep","Revise")'),
    ]
    for k, (label, f) in enumerate(ex):
        ws.cell(row=r0 + 6 + k, column=1, value=label).font = BOLD
        cell = ws.cell(row=r0 + 6 + k, column=2, value=f)
        cell.font = BODY
        cell.alignment = Alignment(horizontal="left")
        if label in ("I-CVI", "Pc", "k*"):
            cell.number_format = "0.00"
    ws.cell(row=r0 + 12, column=2, value=(
        "With five experts every one must rate 3 or 4, so this item is "
        "revised even though its I-CVI is .80.")).font = BODY


wb = Workbook()
read_me(wb)
rating_sheet(
    wb, "Relevance", "Relevance ratings (content validity)",
    "Parts A and B. Part C is background only and is rated for clarity.",
    home + after, ["A. Home", "B. After the session"])
rating_sheet(
    wb, "Clarity", "Clarity ratings",
    "All parts. Low clarity means rewording, not removing.",
    home + after + before,
    ["A. Home", "B. After the session", "C. Before the session"])
yes_no_sheet(
    wb,
    [(n, "Right area?", a, t) for n, _, a, t in home]
    + [(n, "Answer choices OK?", a, t) for n, _, a, t in after + before]
    + [(n, "Keep flag?", q, t) for n, _, q, t in flags])
wb.save(OUT)
print(f"Wrote {OUT}")
